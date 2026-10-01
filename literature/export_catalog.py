"""Export the seven literature categories as GitHub tables and a searchable CSV."""

import argparse
import csv
import hashlib
import html
from pathlib import Path
import re
from urllib.parse import urlsplit, urlunsplit

import openpyxl


CATEGORIES = [
    ("世界模型基础&控制", "01-foundations-control"),
    ("视觉&视频世界模型", "02-visual-video"),
    ("Latent&预测表示", "03-latent-representations"),
    ("结构化物理&多模态", "04-physics-multimodal"),
    ("数据&模拟&跨本体", "05-data-simulation-embodiment"),
    ("世界模型&VLA&规划", "06-policy-planning"),
    ("评估&安全&Benchmark", "07-evaluation-safety"),
]
FIELDS = ["record_id", "category", "source_row", "title", "authors", "year", "venue", "publication_status", "url"]


def text(value):
    return " ".join(str(value if value is not None else "").split())


def public_url(value):
    for match in re.findall(r"https?://[^\s;<>]+", html.unescape(str(value or ""))):
        match = match.rstrip(".,，；。)")
        parts = urlsplit(match)
        host = (parts.hostname or "").lower()
        query = parts.query.lower()
        if not host or parts.username or parts.password:
            continue
        if any(word in host for word in ("vpn.", "localhost", "sciencedirectassets", "shturl.cc")):
            continue
        if re.fullmatch(r"[\d.]+", host) or any(word in query for word in ("token", "signature", "credential", "x-amz-", "expires")):
            continue
        if parts.path.rstrip("/").endswith("/abstract/document"):
            continue
        if host in ("arxiv.org", "www.arxiv.org"):
            identifier = re.search(r"/(?:abs|pdf)/(\d{4}\.\d{4,5})(?:v\d+)?", parts.path)
            if identifier:
                return "https://arxiv.org/abs/" + identifier.group(1)
        return urlunsplit((parts.scheme, parts.netloc, parts.path, parts.query, ""))
    return ""


def markdown(value):
    return html.escape(text(value), quote=False).replace("|", "&#124;").replace("[", "&#91;").replace("]", "&#93;")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    output = Path(__file__).resolve().parent
    pages = output / "categories"
    pages.mkdir(exist_ok=True)
    workbook = openpyxl.load_workbook(args.workbook, data_only=True)
    records, navigation = [], []

    for number, (name, slug) in enumerate(CATEGORIES, 1):
        sheet = workbook[name]
        if not sheet.auto_filter.ref:
            raise ValueError(f"Missing data-table boundary for {name}")
        end = openpyxl.utils.cell.range_boundaries(sheet.auto_filter.ref)[3]
        rows = []
        for source_row in range(2, end + 1):
            values = [sheet.cell(source_row, col).value for col in range(1, 17)]
            if not values[1]:
                continue
            link = public_url(values[6])
            if not link and sheet.cell(source_row, 7).hyperlink:
                link = public_url(sheet.cell(source_row, 7).hyperlink.target)
            record = dict(zip(FIELDS, [f"S{number:02d}-{source_row:04d}", name, source_row, text(values[1]), text(values[2]), text(values[3]), text(values[4]), text(values[5]), link]))
            rows.append(record)
        records.extend(rows)
        navigation.append(f"| [{name}](categories/{slug}.md) | {len(rows)} |")
        lines = [f"# {name}", "", "[返回总目录](../README.md)", "", f"共 {len(rows)} 条分类记录。出版状态沿用协作表，尚未逐条复核。编号对应本次导出快照的 Sheet 与行号。", "", "| 编号 | 年份 | 论文 | 作者/团队 | 来源 | 出版状态 |", "|---|---|---|---|---|---|"]
        for row in rows:
            title = markdown(row["title"])
            if row["url"]:
                title = f"[{title}](<{row['url']}>)"
            lines.append("| " + " | ".join([row["record_id"], markdown(row["year"]), title, markdown(row["authors"]), markdown(row["venue"]), markdown(row["publication_status"])]) + " |")
        (pages / f"{slug}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    with (output / "catalog.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDS)
        writer.writeheader()
        for row in records:
            writer.writerow({key: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value for key, value in row.items()})

    missing = sum(not row["url"] for row in records)
    digest = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    readme = f"""# World Models for Physical AI · Literature Catalog

按七类研究主题整理的候选文献检索目录。快照日期：2026-10-01。

当前收录 **{len(records)} 条分类记录**。跨主题记录保留，因此该数字不是全局唯一论文数，也不是已核验论文数。这里同步标题、作者、年份、来源和链接；实验结论继续在内部协作表中核查。

| 主题 | 分类记录数 |
|---|---:|
{chr(10).join(navigation)}

## 检索和维护

- 在分类页按标题或作者查找；下载 [catalog.csv](catalog.csv) 可按年份、类别、来源筛选。
- `publication_status` 为原表填写，不代表本目录确认接受或发表状态。
- {missing} 条记录暂未导出可用的公开链接；校园代理、临时签名下载链接和不完整链接不进入公开目录。链接能打开也不等于标题与原文已经核验。
- 记录编号用于定位当前快照，不作为论文永久标识。后续核验优先记录 DOI、arXiv ID 和正式出版版本。
- 更新分类与元数据后重新导出：`python export_catalog.py path/to/workbook.xlsx`。依赖：`openpyxl`。

## 本轮新增讨论入口

以下为 2026-10-01 核对的 Astra 相关来源，另列于现有表格快照之外，未计入上述数量：

- GPT 6 Astra as an Embodied Policy [技术报告与代码](https://github.com/anonymous-report-421/GPT-as-Policy)
- An Unexpected Robot Policy [https://arxiv.org/abs/2609.24170](https://arxiv.org/abs/2609.24170)
- Systematically Exploring the Capabilities of GPT-6 Astra as Embodied Policies [https://arxiv.org/abs/2609.38537](https://arxiv.org/abs/2609.38537)
- FLUX 3 Action [技术报告](https://bfl.ai/models/flux-3-action)

前述项目报告与 Galbot 扩展稿中的重合实验应按证据系列处理，不能计作独立重复验证。

## 快照来源

工作簿名：`{args.workbook.name}`。

SHA-256：`{digest}`。

本目录提供文献索引，不分发论文 PDF、实验截图或模型权重。
"""
    (output / "README.md").write_text(readme, encoding="utf-8")
    assert len({row["record_id"] for row in records}) == len(records)
    with (output / "catalog.csv").open(encoding="utf-8-sig", newline="") as file:
        assert sum(1 for _ in csv.DictReader(file)) == len(records)
    print(f"Exported {len(records)} records in {len(CATEGORIES)} categories; {missing} links pending.")


if __name__ == "__main__":
    main()
