"""Checks for lossless list formatting and preservation of catalog notes."""

import contextlib
import csv
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

import export_catalog
from export_catalog import CATEGORIES, markdown, render_category, render_record


class CatalogFormattingTests(unittest.TestCase):
    def record(self, **changes):
        row = dict(record_id="S01-0002", title="Paper [A] | B_*`\\ <tag>",
                   year="2026", authors="Alice; Bob", venue="Original venue",
                   publication_status="Preprint; unverified",
                   url="https://arxiv.org/abs/2601.12345v2")
        row.update(changes)
        return row

    def test_every_field_and_versioned_url_are_preserved(self):
        row = self.record()
        result = "\n".join(render_record(row))
        for key in ("title", "authors", "year", "venue", "publication_status"):
            self.assertIn(markdown(row[key]), result)
        self.assertIn(row["url"], result)
        self.assertIn(row["record_id"], result)
        self.assertIn('id="s01-0002"', result)
        self.assertNotIn("<tag>", result)
        self.assertIn("&#91;A&#93;", result)
        self.assertIn("&#95;&#42;&#96;&#92;", result)

    def test_missing_values_and_placeholders_are_not_fabricated(self):
        row = self.record(record_id="S01-0003", title="【待补】方向", year="",
                          authors="", venue="", publication_status="", url="")
        result = render_category("Test", [self.record(), row])
        self.assertIn("共 2 条冻结分类记录", result)
        self.assertIn("此数量不是唯一论文数或已核验论文数", result)
        self.assertIn("占位项", result)
        self.assertIn("**[年份未填] 【待补】方向**", result)
        self.assertIn("公开链接待补", result)
        self.assertNotIn("[Paper](<>)", result)

    def test_workbook_export_keeps_existing_readme(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            notes = root / "README.md"
            original = b"Manual verification notes\n"
            notes.write_bytes(original)
            workbook = openpyxl.Workbook()
            workbook.remove(workbook.active)
            for name, _ in CATEGORIES:
                sheet = workbook.create_sheet(name)
                sheet.append(["header"] * 16)
                sheet.append([None, "Sample", "Alice", 2026, "arXiv", "Preprint",
                              "https://arxiv.org/abs/2601.12345v2"])
                sheet.auto_filter.ref = "A1:P2"
            path = root / "source.xlsx"
            workbook.save(path)
            with patch.object(export_catalog, "__file__", str(root / "export_catalog.py")), \
                    patch("sys.argv", ["export_catalog.py", str(path)]), \
                    contextlib.redirect_stdout(io.StringIO()):
                export_catalog.main()
            self.assertEqual(notes.read_bytes(), original)
            for _, slug in CATEGORIES:
                self.assertTrue((root / "categories" / f"{slug}.md").is_file())

    def test_checked_in_categories_match_every_frozen_csv_record(self):
        root = Path(__file__).resolve().parent
        with (root / "catalog.csv").open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.DictReader(file))
        rendered_ids = []
        for name, slug in CATEGORIES:
            category_rows = [row for row in rows if row["category"] == name]
            path = root / "categories" / f"{slug}.md"
            with self.subTest(category=name):
                self.assertEqual(path.read_text(encoding="utf-8"),
                                 render_category(name, category_rows))
            rendered_ids.extend(row["record_id"] for row in category_rows)
        self.assertCountEqual(rendered_ids, [row["record_id"] for row in rows])
        self.assertEqual(len(rendered_ids), len(set(rendered_ids)))

    def test_format_only_preserves_csv_notes_and_original_url(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            row = dict(self.record(), category=CATEGORIES[0][0], source_row="2")
            stream = io.StringIO(newline="")
            writer = csv.DictWriter(stream, fieldnames=export_catalog.FIELDS)
            writer.writeheader()
            writer.writerow(row)
            original = stream.getvalue().encode("utf-8-sig")
            (root / "catalog.csv").write_bytes(original)
            notes = b"Existing audit history\n"
            (root / "README.md").write_bytes(notes)
            with patch.object(export_catalog, "__file__", str(root / "export_catalog.py")), \
                    patch("sys.argv", ["export_catalog.py", "--format-only"]), \
                    contextlib.redirect_stdout(io.StringIO()):
                export_catalog.main()
            self.assertEqual((root / "catalog.csv").read_bytes(), original)
            self.assertEqual((root / "README.md").read_bytes(), notes)
            page = root / "categories" / (CATEGORIES[0][1] + ".md")
            self.assertIn(row["url"], page.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
